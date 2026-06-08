import sys
import os
import io
import subprocess
import tempfile  # Для создания временной папки
from PyQt5.QtWidgets import (
    QApplication, QWidget, QPushButton, QLabel,
    QFileDialog, QVBoxLayout
)
from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage  # Переименовали, чтобы не путать с PIL
from PIL import Image as PILImage


def convert_to_docx(input_path):
    if input_path.endswith(".docx"):
        return input_path

    output_dir = os.path.dirname(input_path)

    subprocess.run([
        "soffice",
        "--headless",
        "--convert-to", "docx",
        "--outdir", output_dir,
        input_path
    ])

    new_path = os.path.splitext(input_path)[0] + ".docx"
    return new_path


def extract_content(word_path, excel_path):
    word_path = convert_to_docx(word_path)

    doc = Document(word_path)
    wb = Workbook()
    ws = wb.active

    row_offset = 1

    # --- таблицы ---
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                ws.cell(row=row_offset + i, column=j + 1, value=cell.text)
        row_offset += len(table.rows) + 2

    # --- картинки ---
    img_index = 1

    # Создаем временную папку, которая удалится сама
    with tempfile.TemporaryDirectory() as temp_dir:
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    img_data = rel.target_part.blob
                    image = PILImage.open(io.BytesIO(img_data))

                    # Сохраняем во временную папку
                    temp_path = os.path.join(temp_dir, f"temp_img_{img_index}.png")
                    image.save(temp_path)

                    # Загружаем картинку в openpyxl и добавляем на лист
                    xl_img = OpenpyxlImage(temp_path)

                    # Устанавливаем высоту строки под картинку (примерно 100 пикселей)
                    ws.row_dimensions[row_offset].height = 80

                    # Прикрепляем картинку к левому верхнему углу ячейки A{row_offset}
                    ws.add_image(xl_img, f"A{row_offset}")

                    row_offset += 6  # Шаг по строкам, чтобы картинки не накладывались друг на друга
                    img_index += 1

                except Exception as e:
                    print(f"Пропущено изображение: {e}")
                    continue

        # --- сохранение файла ---
        wb.save(excel_path)


class App(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Word → Excel (любой формат)")

        self.word_path = ""
        self.excel_path = ""

        layout = QVBoxLayout()

        self.label = QLabel("Выбери файл")
        layout.addWidget(self.label)

        btn_word = QPushButton("Выбрать файл")
        btn_word.clicked.connect(self.select_word)
        layout.addWidget(btn_word)

        btn_excel = QPushButton("Сохранить Excel")
        btn_excel.clicked.connect(self.select_excel)
        layout.addWidget(btn_excel)

        btn_run = QPushButton("Запустить")
        btn_run.clicked.connect(self.run)
        layout.addWidget(btn_run)

        self.setLayout(layout)

    def select_word(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Файл", "", "All Files (*);;Word Files (*.docx *.doc)"
        )
        if path:
            self.word_path = path
            self.label.setText(f"Файл: {os.path.basename(path)}")

    def select_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Excel файл", "output.xlsx", "Excel Files (*.xlsx)"
        )
        if path:
            if not path.endswith(".xlsx"):
                path += ".xlsx"
            self.excel_path = path
            self.label.setText(f"Excel: {os.path.basename(path)}")

    def run(self):
        if not self.word_path or not self.excel_path:
            self.label.setText("Выбери файлы!")
            return

        try:
            self.label.setText("Обработка...")
            QApplication.processEvents()
            extract_content(self.word_path, self.excel_path)
            self.label.setText("Готово ✅")
        except Exception as e:
            self.label.setText(f"Ошибка: {e}")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = App()
    window.show()
    sys.exit(app.exec_())
